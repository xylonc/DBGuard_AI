"""Measure: full cell coverage comparison (before/after pruning)."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent / "backend"))

from openpyxl import load_workbook

CIS_PATH = "/workspace/DBGuardAI/local docs/CIS_PostgreSQL_18_Benchmark_v1.0.0.xlsx"

# Count non-empty cells per column across the main sheet
wb = load_workbook(filename=CIS_PATH, read_only=True, data_only=True)
ws = wb['Level 1 - PostgreSQL']
col_non_empty = {}
for row in ws.iter_rows():
    for cell in row:
        col = cell.column  # 1-indexed
        if cell.value is not None:
            col_non_empty[col] = col_non_empty.get(col, 0) + 1
wb.close()

print(f"Non-empty cells per column (Level 1 - PostgreSQL):")
for col in sorted(col_non_empty):
    print(f"  Column {col:2d}: {col_non_empty[col]:4d} non-empty values")
print(f"  Total non-empty cells across all columns: {sum(col_non_empty.values())}")

# Count lines in the current extracted output
xlsx_bytes = open(CIS_PATH, "rb").read()
from app.xlsx_extractor import extract_xlsx_to_text
text = extract_xlsx_to_text(xlsx_bytes)
lines = text.split('\n')

# Check which "kept" columns appear in the output
print(f"\nExtracted text: {len(text):,} chars, {len(lines):,} lines")
print("\nColumns appearing in output:")
for col_num, h in enumerate([
    "Section #", "Recommendation #", "Title", "Assessment Status", "Description",
    "Rationale Statement", "Impact Statement", "Remediation Procedure", "Audit Procedure",
    "Additional Information"
], 1):
    col_lines = [l for l in lines if f"{h}:" in l]
    print(f"  {h:25s}: {len(col_lines):4d} lines")