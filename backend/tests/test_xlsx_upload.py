"""Tests for POST /api/v1/knowledge/upload — XLSX ingestion feature.

Tests:
  A. XLSX extraction
    1. Normal workbook produces meaningful text
    2. Multiple rows are preserved
    3. Multiple sheets are handled
  B. Upload endpoint
    4. Valid XLSX returns success
    5. Non-XLSX file returns 400
    6. Empty workbook returns 400
    7. Corrupt XLSX returns 400
    8. RAG ingestion is called with extracted content
  C. Regression / safety
    9. Long ``# ...`` spreadsheet content cannot cause DB section length overflow
    10. Long cell content is preserved (not silently truncated)
    11. Existing API routes still exist after the change
"""

import io
import sys
from datetime import datetime, timezone
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

# Ensure imports resolve
# Add backend root to path so `app` module is importable
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

# Also add project root for services.rag imports
sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

from openpyxl import Workbook  # type: ignore

from app.main import app as flask_app  # noqa: E402

client = TestClient(flask_app)


# ── helpers ──────────────────────────────────────────────────────────────


def make_xlsx(
    sheet_name: str = "Security Controls",
    header: list[str] | None = None,
    rows: list[list[str]] | None = None,
) -> bytes:
    """Create a tiny XLSX workbook in memory and return its bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    if header:
        ws.append(header)
    if rows:
        for row in rows:
            ws.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    buf.seek(0)
    return buf.read()


def make_empty_xlsx(sheet_name: str = "Empty Sheet") -> bytes:
    """Create an XLSX with a sheet name but no cells at all."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    buf.seek(0)
    return buf.read()


# ── A. XLSX extraction ──────────────────────────────────────────────────


class TestXlsxExtraction:
    """Direct tests on the XLSX extractor module."""

    def test_normal_workbook_produces_meaningful_text(self):
        """A.1 — spreadsheet values are converted to meaningful text."""
        from app.xlsx_extractor import extract_xlsx_to_text

        xlsx_bytes = make_xlsx(
            sheet_name="Security Controls",
            header=["Control", "Requirement", "Severity"],
            rows=[
                ["DB-001", "TLS required", "High"],
                ["DB-002", "Logging enabled", "Medium"],
            ],
        )

        text = extract_xlsx_to_text(xlsx_bytes)

        assert "Security Controls" in text
        assert "DB-001" in text
        assert "TLS required" in text
        assert "DB-002" in text
        assert "Logging enabled" in text
        assert "Control" in text

    def test_multiple_rows_preserved(self):
        """A.2 — multiple rows are preserved in output."""
        from app.xlsx_extractor import extract_xlsx_to_text

        rows = [[f"row-{i}", f"req-{i}", "Low"] for i in range(50)]
        xlsx_bytes = make_xlsx(
            sheet_name="Bulk",
            header=["ID", "Requirement", "Severity"],
            rows=rows,
        )
        text = extract_xlsx_to_text(xlsx_bytes)

        # All 50 rows should appear
        for i in range(50):
            assert f"row-{i}" in text, f"row-{i} missing from extraction"
            assert f"req-{i}" in text, f"req-{i} missing from extraction"

    def test_multiple_sheets_handled(self):
        """A.3 — multiple worksheets are extracted."""
        from openpyxl import Workbook
        from app.xlsx_extractor import extract_xlsx_to_text

        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Sheet One"
        ws1.append(["Col A", "Col B"])
        ws1.append(["val1", "val2"])

        ws2 = wb.create_sheet("Sheet Two")
        ws2.append(["Col X", "Col Y"])
        ws2.append(["valX", "valY"])

        buf = io.BytesIO()
        wb.save(buf)
        xlsx_bytes = buf.getvalue()
        wb.close()

        text = extract_xlsx_to_text(xlsx_bytes)

        assert "Sheet One" in text
        assert "Sheet Two" in text
        assert "val1" in text
        assert "valX" in text


# ── B. Upload endpoint ──────────────────────────────────────────────────


class TestXlsxUploadEndpoint:
    """Tests for the POST /api/v1/knowledge/upload endpoint."""

    def test_valid_xlsx_upload_succeeds(self, mocker):
        """B.4 — valid XLSX uploads and reaches existing RAG ingestion path."""
        mock_result = mocker.MagicMock()
        mock_result.status = "ingested"
        mock_result.chunks_created = 5
        mock_result.errors = []

        mocker.patch(
            "app.main.RAGService.ingest_document",
            return_value=mock_result,
        )

        xlsx_bytes = make_xlsx(
            sheet_name="Security Controls",
            header=["Control", "Requirement"],
            rows=[
                ["DB-001", "TLS required"],
                ["DB-002", "Logging enabled"],
            ],
        )

        response = client.post(
            "/api/v1/knowledge/upload",
            files={
                "file": (
                    "security-controls.xlsx",
                    io.BytesIO(xlsx_bytes),
                    "application/octet-stream",
                )
            },
        )

        assert response.status_code == 200
        data = response.json()
        assert data["status"] == "ingested"
        assert data["chunks_created"] == 5
        assert "document_id" in data
        assert "title" in data

    def test_non_xlsx_file_rejected(self):
        """B.5 — a non-XLSX file is rejected with 400."""
        response = client.post(
            "/api/v1/knowledge/upload",
            files={
                "file": ("readme.txt", io.BytesIO(b"not an xlsx"), "text/plain")
            },
        )

        assert response.status_code == 400
        assert ".xlsx" in response.json()["detail"].lower()

    def test_empty_xlsx_rejected(self):
        """B.6 — an empty workbook (no cells) is rejected with 400."""
        xlsx_bytes = make_empty_xlsx()

        response = client.post(
            "/api/v1/knowledge/upload",
            files={
                "file": (
                    "empty.xlsx",
                    io.BytesIO(xlsx_bytes),
                    "application/octet-stream",
                )
            },
        )

        # Empty workbooks produce short content that the RAG service rejects.
        # The extractor detects it first (400) or the RAG service rejects it (400).
        # In either case, it should NOT be a success.
        assert response.status_code == 400, (
            f"Expected 400 for empty workbook, got {response.status_code}"
        )

    def test_empty_file_rejected(self):
        """B.6b — an uploaded file with zero bytes is rejected."""
        response = client.post(
            "/api/v1/knowledge/upload",
            files={
                "file": (
                    "empty.xlsx",
                    io.BytesIO(b""),
                    "application/octet-stream",
                )
            },
        )

        assert response.status_code == 400

    def test_corrupt_xlsx_rejected(self):
        """B.7 — a corrupt/invalid XLSX is rejected with 400."""
        # Valid XLSX header followed by garbage bytes
        fake_xlsx = b"PK\x03\x04" + b"\x00" * 100

        response = client.post(
            "/api/v1/knowledge/upload",
            files={
                "file": (
                    "corrupt.xlsx",
                    io.BytesIO(fake_xlsx),
                    "application/octet-stream",
                )
            },
        )

        # Must be 400 (parse error) - corrupt input is a client error
        assert response.status_code == 400, (
            f"Expected 400 for corrupt XLSX, got {response.status_code}"
        )

    def test_rag_ingestion_called_with_extracted_content(self, mocker):
        """B.8 — RAG ingestion is called with the extracted text."""
        captured_doc = {}

        def capture_ingest(doc):
            captured_doc["document_id"] = doc.document_id
            captured_doc["content"] = doc.content
            captured_doc["title"] = doc.title
            mock_result = mocker.MagicMock()
            mock_result.status = "ingested"
            mock_result.chunks_created = 3
            mock_result.errors = []
            return mock_result

        mocker.patch(
            "app.main.RAGService.ingest_document",
            side_effect=capture_ingest,
        )

        xlsx_bytes = make_xlsx(
            sheet_name="Audit Findings",
            header=["ID", "Finding", "Details"],
            rows=[
                ["F-001", "Weak encryption detected", "TLS 1.0 is enabled on port 443 and must be disabled according to CIS Benchmark section 4.2.1"],
                ["F-002", "Insufficient logging", "Audit logging is disabled for database operations on the production server which violates security policy requirements"],
                ["F-003", "Missing firewall rules", "The network firewall configuration does not include any deny-all rules for the database subnet and allows inbound traffic from any source"],
            ],
        )

        client.post(
            "/api/v1/knowledge/upload",
            files={
                "file": (
                    "audit.xlsx",
                    io.BytesIO(xlsx_bytes),
                    "application/octet-stream",
                )
            },
        )

        assert "content" in captured_doc
        assert "Weak encryption detected" in captured_doc["content"]
        assert "F-001" in captured_doc["content"]
        assert captured_doc["title"].lower().startswith("audit")


# ── C. Regression / safety ──────────────────────────────────────────────


def make_long_header_xlsx() -> bytes:
    """Build an XLSX whose first data row starts with ``# `` and has a long body.

    This reproduces the bug where a worksheet header cell like
    ``# PostgreSQL CIS Benchmark v1.0.0`` causes the extractor to produce
    lines such as::

        # PostgreSQL CIS Benchmark v1.0.0: Enable detailed query logging …

    The chunker then treats the entire rest of the line as a section header,
    producing section values that exceed ``VARCHAR(255)`` and fail the INSERT.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Benchmarks"
    ws.append(["# PostgreSQL CIS Benchmark v1.0.0"])
    ws.append(["Section 1: Installation and Configuration"])
    ws.append([
        "Enable detailed query logging for audit trail purposes on the "
        "PostgreSQL database server instance configured to capture all "
        "connection attempts and SQL queries executed by user roles across "
        "all databases in the production environment with strict compliance "
        "requirements for financial institutions operating in multi-tenant "
        "cloud environments with shared infrastructure and regulatory "
        "oversight bodies requiring comprehensive audit capabilities"
    ])
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    buf.seek(0)
    return buf.getvalue()


class TestChunkSectionBounded:
    """Regression tests: chunk section values must never exceed 255 chars."""

    def test_xlsx_with_long_hash_header_section_is_bounded(self):
        """C.9 — XLSX containing a long ``# `` header line must not produce a
        ``section`` value longer than 255 characters.
        """
        from app.xlsx_extractor import extract_xlsx_to_text
        from services.rag.rag_service import RAGService, KnowledgeDocument

        xlsx_bytes = make_long_header_xlsx()
        text = extract_xlsx_to_text(xlsx_bytes)

        # The extractor escapes # lines as \# lines.
        # The key assertion: the long content is preserved and doesn't cause
        # any chunk section to exceed 255 chars.
        assert "CIS Benchmark" in text, "Test setup: extracted text must contain the long header line"
        assert "\\# PostgreSQL" in text  # Extracted as escaped # line

        doc = KnowledgeDocument(
            document_id="regression-xlsx-long-header",
            title="Benchmarks",
            version="1.0.0",
            content=text,
            effective_date=datetime.now(timezone.utc),
            status="draft",
        )

        chunks = RAGService()._chunk_document(doc)

        assert len(chunks) > 0, "Expected at least one chunk"

        for i, chunk in enumerate(chunks):
            assert len(chunk.section) <= 255, (
                f"Chunk {i} section is {len(chunk.section)} chars "
                f"(exceeds VARCHAR(255)): {chunk.section[:60]}..."
            )

    def test_direct_300_char_section_header_is_clamped(self):
        """C.9b — A document with a direct 300-char ``#`` header must clamp to 255."""
        from services.rag.rag_service import RAGService, KnowledgeDocument

        content = f"{'#' + 'X' * 300}\nBody text\n"
        doc = KnowledgeDocument(
            document_id="regression-300-char-header",
            title="Test",
            version="1.0.0",
            content=content,
            effective_date=datetime.now(timezone.utc),
            status="draft",
        )

        chunks = RAGService()._chunk_document(doc)
        assert len(chunks) > 0

        for i, chunk in enumerate(chunks):
            assert len(chunk.section) <= 255, (
                f"Chunk {i} section is {len(chunk.section)} chars (should be clamped)"
            )

    def test_split_oversized_chunk_section_stays_within_255(self):
        """C.10 — When an oversized chunk is split, the ``(cont.)`` suffix must
        not push the section beyond 255 characters.
        """
        from services.rag.rag_service import RAGService, KnowledgeChunk

        long_section = "Y" * 255
        long_content = "Z" * 5000

        chunk = KnowledgeChunk(
            document_id="regression-split",
            section=long_section,
            content=long_content,
            chunk_hash="abc",
            chunk_index=0,
            postgresql_versions=["15"],
            environment_applicability=["all"],
            source_document_title="Test",
            source_document_version="1.0",
        )

        sub_chunks = RAGService()._split_oversized_chunk(chunk)

        assert len(sub_chunks) > 1, "Expected at least 2 sub-chunks from 5000-char content"

        for i, sc in enumerate(sub_chunks):
            assert len(sc.section) <= 255, (
                f"Split chunk {i} section is {len(sc.section)} chars (with '(cont.)')"
            )
            if i > 0:
                assert sc.section.endswith(" (cont.)"), (
                    f"Split chunk {i} section should end with '(cont.)'"
                )

    def test_long_xlsx_content_preserved_in_chunk_text(self):
        """C.10b — The full long spreadsheet text must remain in the chunk ``content``
        (``TEXT`` column) even though ``section`` is bounded.
        """
        from app.xlsx_extractor import extract_xlsx_to_text
        from services.rag.rag_service import RAGService, KnowledgeDocument

        xlsx_bytes = make_long_header_xlsx()
        text = extract_xlsx_to_text(xlsx_bytes)

        doc = KnowledgeDocument(
            document_id="regression-preserve-content",
            title="Benchmarks",
            version="1.0.0",
            content=text,
            effective_date=datetime.now(timezone.utc),
            status="draft",
        )

        chunks = RAGService()._chunk_document(doc)
        combined_content = "".join(c.content for c in chunks)

        assert "Enable detailed query logging for audit trail purposes" in combined_content
        assert "shared infrastructure" in combined_content
        assert "regulatory oversight bodies" in combined_content

    def test_shell_code_lines_not_fake_sections(self):
        """C.10c — Lines beginning with ``# `` inside XLSX cells must NOT be
        interpreted as Markdown section headers by the RAG chunker.
        """
        from app.xlsx_extractor import extract_xlsx_to_text
        from services.rag.rag_service import RAGService, KnowledgeDocument

        wb = Workbook()
        ws = wb.active
        ws.title = "Code Tests"
        ws.append(["Section #", "Recommendation #", "Title", "Remediation Procedure"])
        ws.append([
            "1", "1.1", "Test",
            "# psql\n# whoami\n# chmod 600 /etc/config\n# systemctl restart",
        ])

        buf = io.BytesIO()
        wb.save(buf)
        xlsx_bytes = buf.getvalue()

        text = extract_xlsx_to_text(xlsx_bytes)
        doc = KnowledgeDocument(
            document_id="hash-guard-test",
            title="Hash guard",
            version="1.0",
            content=text,
            effective_date=datetime.now(timezone.utc),
            status="draft",
        )

        chunks = RAGService()._chunk_document(doc)
        all_text = " ".join(c.content for c in chunks)

        # The # guarded lines should be present as content
        assert "# psql" in all_text
        assert "# whoami" in all_text
        # But the chunks must NOT have section names like "psql", "whoami" etc.
        section_names = [c.section for c in chunks]
        for section in section_names:
            assert "psql" not in section.lower(), (
                f"Chunk section '{section}' contains '# ' guard violation"
            )
            assert "whoami" not in section.lower(), (
                f"Chunk section '{section}' contains '# ' guard violation"
            )

    def test_long_text_preserved_in_output(self):
        """C.10d — Long cell values must survive normalization and chunking intact.
        A multi-paragraph, >1000-character value should be fully retrievable.
        """
        from app.xlsx_extractor import extract_xlsx_to_text
        from services.rag.rag_service import RAGService, KnowledgeDocument

        long_text = " ".join(
            f"This is paragraph {i}. It contains important audit information "
            f"about control implementation details and security posture "
            f"assessment for compliance evidence purposes. "
            f"Line {i} of the remediation procedure."
            for i in range(1, 30)
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "Long Text Test"
        ws.append(["Section #", "Recommendation #", "Title", "Remediation Procedure"])
        ws.append(["1", "1.1", "Long remediation", long_text])

        buf = io.BytesIO()
        wb.save(buf)
        xlsx_bytes = buf.getvalue()

        text = extract_xlsx_to_text(xlsx_bytes)
        doc = KnowledgeDocument(
            document_id="long-text-test",
            title="Long text",
            version="1.0",
            content=text,
            effective_date=datetime.now(timezone.utc),
            status="draft",
        )

        chunks = RAGService()._chunk_document(doc)
        all_text = " ".join(c.content for c in chunks)

        for i in range(1, 30):
            assert f"This is paragraph {i}." in all_text, (
                f"Paragraph {i} was lost during normalization/chunking"
            )


class TestRegressionRoutes:
    """C.11 — Existing API routes still exist after the XLSX change."""

    def test_proposals_compile_route_exists(self):
        route_paths = [route.path for route in flask_app.routes]
        assert "/api/v1/proposals/compile" in route_paths

    def test_templates_ingest_route_exists(self):
        route_paths = [route.path for route in flask_app.routes]
        assert "/api/v1/templates/ingest" in route_paths
        assert "/api/v1/templates/search" in route_paths

    def test_knowledge_routes_exist(self):
        route_paths = [route.path for route in flask_app.routes]
        assert "/api/v1/knowledge/documents" in route_paths
        assert "/api/v1/knowledge/search" in route_paths

    def test_xlsx_upload_route_exists(self):
        route_paths = [route.path for route in flask_app.routes]
        assert "/api/v1/knowledge/upload" in route_paths

    def test_health_route_exists(self):
        route_paths = [route.path for route in flask_app.routes]
        assert "/api/v1/health" in route_paths
