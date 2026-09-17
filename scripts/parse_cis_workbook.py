#!/usr/bin/env python3
"""Parse CIS PostgreSQL 17 Benchmark workbook into records JSON."""

import argparse
import hashlib
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


class WorkbookFormatError(Exception):
    """Raised when the workbook format is invalid."""
    pass


def parse_workbook(input_path: str, output_path: str) -> None:
    """Parse the CIS workbook and output records JSON."""
    wb = load_workbook(input_path, data_only=True)
    
    # Get the Combined Profiles sheet
    if "Combined Profiles" not in wb.sheetnames:
        raise WorkbookFormatError("Missing required sheet: 'Combined Profiles'")
    
    sheet = wb["Combined Profiles"]
    
    # Read header row and find column indices by name
    headers = [cell.value for cell in sheet[1]]
    
    required_columns = [
        "Recommendation #",
        "Section #",
        "Profile",
        "Title",
        "Assessment Status",
        "Description",
        "Rationale Statement",
        "Impact Statement",
        "Audit Procedure",
        "Remediation Procedure",
        "Additional Information",
        "References",
        "Default Value",
    ]
    
    # Check for missing headers
    missing = [col for col in required_columns if col not in headers]
    if missing:
        raise WorkbookFormatError(f"Missing required header(s): {', '.join(missing)}")
    
    # Build column index map
    col_idx = {col: headers.index(col) for col in required_columns}
    
    # Parse records
    records = []
    seen_recommendations = set()
    
    for row in range(2, sheet.max_row + 1):
        recommendation = sheet.cell(row, col_idx["Recommendation #"] + 1).value
        
        # Skip section rows (empty Recommendation #)
        if recommendation is None or str(recommendation).strip() == "":
            continue
        
        # Check for duplicate recommendation
        if recommendation in seen_recommendations:
            raise WorkbookFormatError(f"Duplicate recommendation: {recommendation}")
        seen_recommendations.add(recommendation)
        
        # Extract cell values (empty cell -> null)
        def get_cell(col_name):
            cell = sheet.cell(row, col_idx[col_name] + 1).value
            return cell if cell is not None else None
        
        record = {
            "recommendation": get_cell("Recommendation #"),
            "section": get_cell("Section #"),
            "profile": get_cell("Profile"),
            "title": get_cell("Title"),
            "assessment_status": get_cell("Assessment Status"),
            "description": get_cell("Description"),
            "rationale": get_cell("Rationale Statement"),
            "impact": get_cell("Impact Statement"),
            "audit_procedure": get_cell("Audit Procedure"),
            "remediation_procedure": get_cell("Remediation Procedure"),
            "additional_information": get_cell("Additional Information"),
            "references": get_cell("References"),
            "pg_default_value": get_cell("Default Value"),
            "source_row": row,
        }
        
        # Compute source_sha256
        sha_data = {
            "recommendation": record["recommendation"],
            "title": record["title"],
            "assessment_status": record["assessment_status"],
            "audit_procedure": record["audit_procedure"],
            "remediation_procedure": record["remediation_procedure"],
        }
        sha_str = json.dumps(sha_data, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
        record["source_sha256"] = hashlib.sha256(sha_str.encode("utf-8")).hexdigest()
        
        records.append(record)
    
    # Compute source file hash
    with open(input_path, "rb") as f:
        source_sha256 = hashlib.sha256(f.read()).hexdigest()
    
    # Build output structure
    output_data = {
        "benchmark": "CIS PostgreSQL 17 Benchmark",
        "benchmark_version": "1.1.0",
        "sheet": "Combined Profiles",
        "source_file_sha256": source_sha256,
        "record_count": len(records),
        "records": records,
    }
    
    # Ensure output directory exists
    output_dir = Path(output_path).parent
    output_dir.mkdir(parents=True, exist_ok=True)
    
    # Write output file
    with open(output_path, "w", encoding="utf-8", newline="\n") as f:
        json.dump(output_data, f, sort_keys=True, indent=2, ensure_ascii=False)
        f.write("\n")
    
    # Print summary
    print(f"Record count: {len(records)}")
    
    # Count by (profile, assessment_status)
    counts = Counter((r["profile"], r["assessment_status"]) for r in records)
    print("\nCounts by (profile, assessment_status):")
    for (profile, status), count in sorted(counts.items()):
        print(f"  ({profile!r}, {status!r}): {count}")


def main():
    parser = argparse.ArgumentParser(description="Parse CIS PostgreSQL 17 Benchmark workbook")
    parser.add_argument(
        "--input",
        default="local docs/CIS_PostgreSQL_17_Benchmark_v1.1.0.xlsx",
        help="Input XLSX file path",
    )
    parser.add_argument(
        "--output",
        default="catalog/benchmarks/cis-pg17-v1.1.0/records.json",
        help="Output JSON file path",
    )
    args = parser.parse_args()
    
    try:
        parse_workbook(args.input, args.output)
    except WorkbookFormatError as e:
        print(str(e), file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
